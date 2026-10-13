"""GD2 baseline/event matching, read-only. No Apps Script or Sheet writes.

Usage: python 15_check_legal_effect.py --mode baseline --workbook backup.xlsx
       python 15_check_legal_effect.py --mode events --workbook backup.xlsx --step08-json output/vbqppl_nhap_batch_payload.json
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from legal_effect_analysis import normalize_number
from legal_effect_identity import inspect_catalog
from legal_effect_official import check_official_sources, load_official_sources

HEADERS = [
    "Mã phiếu rà soát", "Khóa văn bản chuẩn", "Số hiệu", "Loại văn bản",
    "Tên văn bản", "Nguồn xuất hiện", "ID nguồn VBQPPL", "ID nguồn VBQPPL_Nhap",
    "URL văn bản gốc", "Ngày bắt đầu hiệu lực", "Tình trạng hiệu lực đề xuất",
    "Loại quan hệ", "Văn bản tác động", "Ngày hiệu lực tác động", "Phạm vi tác động",
    "Trích dẫn điều khoản", "URL nguồn chính thức", "URL nguồn bổ trợ",
    "Mức tin cậy", "Trạng thái xác minh", "Người xác minh", "Thời gian xác minh",
    "Kết luận pháp chế", "Thời gian quét", "Khóa sự kiện", "Dấu vân tay nguồn",
    "Lỗi hoặc ngoại lệ", "Ghi chú vận hành",
]
SOURCES = ("VBQPPL", "VBQPPL_Nhap")


def catalog_from_workbook(path: Path) -> dict[str, list[dict]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    result: dict[str, list[dict]] = defaultdict(list)
    try:
        for sheet in SOURCES:
            if sheet not in wb.sheetnames:
                raise ValueError("Missing source sheet: " + sheet)
            cells = iter(wb[sheet].iter_rows(values_only=True))
            headers = [str(x or "").strip() for x in next(cells)]
            for physical_row, values in enumerate(cells, start=2):
                row = dict(zip(headers, values))
                code = normalize_number(row.get("Số hiệu"))
                if not code:
                    continue
                result[code].append({
                    "sheet": sheet,
                    "row_number": physical_row,
                    "id": str(row.get("ID VĂN BẢN") or row.get("ID Văn bản") or ""),
                    "type": str(row.get("Loại văn bản") or ""),
                    "title": str(row.get("Tên văn bản") or ""),
                    "url": str(row.get("Link Văn bản") or ""),
                    "effect_start": str(row.get("Ngày hiệu lực") or ""),
                })
    finally:
        wb.close()
    return dict(result)


def build_review(catalog: dict, events: list[dict], mode: str) -> dict:
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    identity = inspect_catalog(catalog)
    candidates = [None] * len(catalog) if mode == "baseline" else events
    if mode == "baseline":
        pairs = [(code, None) for code in sorted(catalog)]
    else:
        pairs = [(normalize_number(x.get("old_number")), x) for x in candidates if isinstance(x, dict)]
    rows, unmatched = {}, []
    for code, event in pairs:
        if code not in catalog:
            unmatched.append({"old_number": code, "event_key": (event or {}).get("event_key", ""), "reason": "NOT_IN_CATALOG"})
            continue
        sources = catalog[code]
        identity_row = identity["by_number"][code]
        representative = identity_row["representative"]
        issues = identity_row["issues"]
        collision = identity_row["requires_identity_review"]
        source_url = identity_row["preferred_url"]
        event = event or {}
        event_key = event.get("event_key", "")
        seed = code + "|" + (event_key or "BASELINE")
        review_id = "RV-" + hashlib.sha256(seed.encode()).hexdigest()[:16].upper()
        values = [
            review_id, code, code, representative["type"], representative["title"],
            ", ".join(sorted({s["sheet"] for s in sources})),
            "; ".join(sorted({s["id"] for s in sources if s["sheet"] == "VBQPPL" and s["id"]})),
            "; ".join(sorted({s["id"] for s in sources if s["sheet"] == "VBQPPL_Nhap" and s["id"]})),
            source_url, representative["effect_start"],
            "CHUA_XAC_MINH", event.get("relation_type", ""), event.get("new_number", ""),
            event.get("effect_date_candidate", ""), event.get("affected_scope", ""),
            event.get("evidence_text", ""), "", event.get("source_url", ""),
            event.get("confidence", "THAP"),
            "CAN_BO_SUNG_CHUNG_CU" if collision else "CHO_XAC_MINH",
            "", "", "", now, event_key,
            identity_row["fingerprint"],
            "; ".join(issues),
            event.get("note", "Rà soát ban đầu; chưa có kết luận pháp lý"),
        ]
        rows[review_id] = dict(zip(HEADERS, values))
    return {
        "generated_at": now, "mode": mode, "read_only": True,
        "summary": {
            "catalog_unique_numbers": len(catalog),
            "catalog_records": sum(map(len, catalog.values())),
            "duplicated_numbers": sum(len(v) > 1 for v in catalog.values()),
            "events_received": len(events), "review_rows": len(rows),
            "unmatched_events": len(unmatched), "pending_verification": len(rows),
        },
        "records_review": [rows[k] for k in sorted(rows)],
        "unmatched_events": unmatched,
        "identity_audit": {"summary": identity["summary"], "issues": identity["issues"]},
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mode", choices=("baseline", "events"), required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--step08-json", type=Path)
    parser.add_argument("--official-index", type=Path, help="JSON mapping document number to verified official full-text URL")
    parser.add_argument("--output-dir", type=Path, default=Path("output"))
    args = parser.parse_args(argv)
    if args.mode == "events" and args.step08_json is None:
        parser.error("--mode events requires --step08-json")
    events = []
    if args.step08_json:
        source = json.loads(args.step08_json.read_text(encoding="utf-8-sig"))
        if "legal_effect_candidates" not in source:
            parser.error("Missing legal_effect_candidates: run updated step08")
        events = source["legal_effect_candidates"]
        if not isinstance(events, list):
            parser.error("legal_effect_candidates must be a list")
    report = build_review(catalog_from_workbook(args.workbook), events, args.mode)
    if args.official_index:
        if args.mode != "events":
            parser.error("--official-index requires --mode events")
        index = load_official_sources(args.official_index)
        report["official_verification"] = check_official_sources(report, events, index)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "legal_effect_review.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    with (args.output_dir / "legal_effect_review.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(report["records_review"])
    print(json.dumps(report["summary"], ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
